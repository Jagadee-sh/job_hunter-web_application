from pathlib import Path

from docx import Document
from reportlab.lib.pagesizes import LETTER
from reportlab.pdfgen.canvas import Canvas


class ResumeEngine:
    def tailor(self, base_path: Path, output_docx: Path, output_pdf: Path, keywords: list[str]) -> None:
        document = Document(str(base_path))
        if keywords:
            paragraph = document.add_paragraph()
            paragraph.add_run("Relevant skills: ").bold = True
            paragraph.add_run(", ".join(dict.fromkeys(keywords)))
        document.save(str(output_docx))
        canvas = Canvas(str(output_pdf), pagesize=LETTER)
        y = 750
        for paragraph in document.paragraphs:
            if paragraph.text.strip():
                canvas.drawString(54, y, paragraph.text[:120])
                y -= 14
                if y < 54:
                    canvas.showPage()
                    y = 750
        canvas.save()
